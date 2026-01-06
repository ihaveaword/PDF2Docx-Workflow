import os
import re
from pdf2image import convert_from_path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

# === 配置区域 ===
PDF_FOLDER = 'PDF_Source'      # 你的PDF文件夹
OUTPUT_FOLDER = 'Output_Docs'  # 结果输出文件夹

# 确保输出目录存在
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

def setup_page_layout(doc):
    """设置页面为 A4 横向，并应用自定义页边距"""
    section = doc.sections[0]
    
    # 1. 设置方向为横向 (Landscape)
    section.orientation = WD_ORIENT.LANDSCAPE
    # 必须手动交换宽高才能真正生效
    section.page_width, section.page_height = section.page_height, section.page_width
    
    # 2. 设置页边距 (根据你的截图)
    section.top_margin = Cm(0.6)
    section.bottom_margin = Cm(0.6)
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)

def extract_sort_id(filename):
    """
    从文件名中提取序号。
    例如从 "授权专利-发明专利 9-设备管理..." 中提取出 9
    逻辑：提取文件名中出现的第一个独立数字
    """
    match = re.search(r'(\d+)', filename)
    if match:
        return int(match.group(1))
    return 9999 # 如果没找到数字，排在最后

def process_files(max_count=None):
    # 获取所有PDF文件
    files = [f for f in os.listdir(PDF_FOLDER) if f.lower().endswith('.pdf')]
    
    # 关键步骤：按文件名里的数字排序 (保证 1, 2, ... 9, 10 的顺序)
    files.sort(key=extract_sort_id)
    
    # 如果指定了处理数量，进行切片 (自动处理 100 > 82 的情况)
    if max_count:
        files = files[:max_count]
    
    total_to_process = len(files)
    print(f"🚀 开始处理前 {total_to_process} 个PDF文件...")

    for index, filename in enumerate(files):
        pdf_path = os.path.join(PDF_FOLDER, filename)
        
        # 提取序号用于 "附件 2-X"
        seq_id = extract_sort_id(filename)
        
        # 去掉后缀的文件名
        file_title = os.path.splitext(filename)[0]
        
        print(f"[{index+1}/{total_to_process}] 正在处理: 附件 2-{seq_id}")

        doc = Document()
        setup_page_layout(doc)
        
        # --- 第一部分：文字头部 ---
        p = doc.add_paragraph()
        # 格式：附件 + 文件名
        run = p.add_run(f"附件 {file_title}")
        run.font.name = '等线'  # 设置西文字体
        # 设置中文字体（东亚字体）
        run._element.rPr.rFonts.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '等线')
        run.font.size = Pt(10.5)  # 五号字 = 10.5磅
        run.bold = True  # 加粗
        
        # 回车 (空行) - 设置字体避免默认明朝
        empty_p = doc.add_paragraph()
        empty_run = empty_p.add_run()
        empty_run.font.name = '等线'
        empty_run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '等线')
        # doc.add_paragraph() # 如果觉得空隙不够，可以解开这行再加一个回车

        # --- 第二部分：左右双图布局 ---
        # 创建一个 1行 2列 的表格
        table = doc.add_table(rows=1, cols=2)
        
        # 设置表格宽度撑满页面 (可选，python-docx默认通常会自适应)
        table.autofit = True 
        
        # 获取左右单元格
        cell_left = table.cell(0, 0)
        cell_right = table.cell(0, 1)

        # === 左侧：插入 PDF 截图 ===
        try:
            # 这里的 Paragraph 是为了让图片居中或对齐，直接 add_picture 到 cell 也可以
            p_left = cell_left.paragraphs[0]
            p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # PDF 转图片
            images = convert_from_path(pdf_path, first_page=1, last_page=1)
            temp_img = f"temp_{seq_id}.jpg"
            images[0].save(temp_img, 'JPEG')
            
            # 插入图片，宽度设为 11厘米 (防止分页)
            run_left = p_left.add_run()
            run_left.add_picture(temp_img, width=Cm(11))
            
            # 清理临时图片
            if os.path.exists(temp_img):
                os.remove(temp_img)
                
        except Exception as e:
            cell_left.text = f"PDF读取失败: {e}"
            print(f"   ❌ PDF错误: {e}")

        # === 右侧：插入占位符 (等待你手动粘贴网页截图) ===
        p_right = cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_right = p_right.add_run("[ 此处请粘贴网页查询截图 ]")
        run_right.font.name = '等线'  # 设置西文字体
        run_right._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '等线')  # 设置中文字体
        run_right.font.size = Pt(14)
        run_right.font.color.rgb = RGBColor(200, 200, 200) # 浅灰色提示文字

        # 保存文件
        # 文件名格式：序号. 原文件名.docx
        save_name = f"{seq_id}. {file_title}.docx"
        doc.save(os.path.join(OUTPUT_FOLDER, save_name))

    print(f"\n🎉 处理完毕！共生成 {total_to_process} 个文件。请检查 Output_Docs 文件夹。")

if __name__ == '__main__':
    # 在这里修改你想处理的文件数量 (比如 2 或 100)
    # 如果想处理全部，可以写 None 或者一个很大的数字
    process_files(max_count=None)  # 处理全部文件